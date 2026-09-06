from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document


def inspect_docx(path: Path) -> dict:
    doc = Document(path)
    paragraphs = []
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            paragraphs.append(
                {
                    "index": index,
                    "style": paragraph.style.name if paragraph.style else None,
                    "text": text,
                }
            )

    tables = []
    for table_index, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({"index": table_index, "rows": rows})

    sections = []
    for index, section in enumerate(doc.sections):
        sections.append(
            {
                "index": index,
                "width": section.page_width,
                "height": section.page_height,
                "top_margin": section.top_margin,
                "bottom_margin": section.bottom_margin,
                "left_margin": section.left_margin,
                "right_margin": section.right_margin,
            }
        )

    return {
        "path": str(path),
        "paragraph_count": len(doc.paragraphs),
        "nonempty_paragraphs": paragraphs,
        "table_count": len(doc.tables),
        "tables": tables,
        "inline_shape_count": len(doc.inline_shapes),
        "section_count": len(doc.sections),
        "sections": sections,
        "core_properties": {
            "title": doc.core_properties.title,
            "subject": doc.core_properties.subject,
            "author": doc.core_properties.author,
            "last_modified_by": doc.core_properties.last_modified_by,
            "created": doc.core_properties.created.isoformat() if doc.core_properties.created else None,
            "modified": doc.core_properties.modified.isoformat() if doc.core_properties.modified else None,
        },
    }


if __name__ == "__main__":
    source = Path(sys.argv[1])
    destination = Path(sys.argv[2])
    destination.write_text(
        json.dumps(inspect_docx(source), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
