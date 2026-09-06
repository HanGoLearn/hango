from __future__ import annotations

import copy
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def set_paragraph_text(paragraph, text: str) -> None:
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    paragraph.add_run(text)


def replace_exact(doc: Document, old: str, new: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == old:
            set_paragraph_text(paragraph, new)
            return
    raise ValueError(f"Paragraph not found: {old[:80]}")


def insert_before(paragraph, text: str, style: str) -> None:
    new_paragraph = paragraph.insert_paragraph_before(text)
    new_paragraph.style = style


def set_cell_text(cell, text: str) -> None:
    if not cell.paragraphs:
        cell.add_paragraph()
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def append_table_row(table, values: list[str], template_index: int = -1) -> None:
    source = table.rows[template_index]._tr
    cloned = copy.deepcopy(source)
    table._tbl.append(cloned)
    row = table.rows[-1]
    for cell, value in zip(row.cells, values, strict=True):
        set_cell_text(cell, value)


def update_table_cell(table, row: int, col: int, text: str) -> None:
    set_cell_text(table.rows[row].cells[col], text)


def main(source: Path, destination: Path) -> None:
    doc = Document(source)

    replacements = {
        "VERSIÓN DE TRABAJO · 2.1 · 31 DE AGOSTO DE 2026":
            "VERSIÓN DE TRABAJO · 2.2 · 5 DE SEPTIEMBRE DE 2026",
        "TESIS  HanGo es un ecosistema de aprendizaje de mandarín que busca convertir situaciones reales, objetivos de clase y datos del alumno en práctica guiada, conversación, feedback y progreso observable.":
            "TESIS  HanGo es un ecosistema de aprendizaje de mandarín que acompaña al estudiante desde sus primeros pasos hasta el perfeccionamiento. Convierte situaciones reales, objetivos de aprendizaje y datos del alumno en práctica guiada, conversación, feedback y progreso observable.",
        "HanGo ya cuenta con un producto local amplio. El núcleo combina una ruta de aprendizaje, orientación cultural para principiantes, placement integrado, ejercicios, comprensión, speaking en tiempo real, writing, reflexión, progreso y un tutor de IA llamado Han. También incorpora una capa de motivación con avatares, HanCoins, logros, tienda y leaderboard. Sin embargo, amplitud funcional no equivale todavía a producto validado: falta llevar una versión estable a producción, seleccionar una ruta pequeña, revisar su mandarín con criterio humano y comprobar que los estudiantes regresan porque aprenden y reciben valor.":
            "HanGo ya cuenta con un producto local amplio para principiantes absolutos y estudiantes que buscan perfeccionar su chino. El núcleo combina Learning Journey, fundamentos con pinyin y tonos, placement integrado, ejercicios, comprensión, Speaking de lección, Speaking Lab, Writing, Reflection, Progress y Han Tutor. El espacio de Han añade Han Companion para conversación por voz y Han Vocabulary para consulta contextual. La motivación incluye avatares, HanCoins, logros, tienda y leaderboard. Esta amplitud todavía debe validarse con usuarios: falta llevar una versión estable a producción, seleccionar una ruta piloto, revisar su mandarín con criterio humano y comprobar aprendizaje, retorno y costo sostenible.",
        "La decisión estratégica sigue siendo concentrar el MVP únicamente en HanGo Learn. La personalización y HanCoins ya forman parte experimental de Learn; no constituyen una línea separada ni sustituyen el valor educativo. Teacher, Institutions, Connect, Community, Play y Marketplace continúan como visión de largo plazo y no deben competir con la validación del núcleo.":
            "La decisión estratégica sigue siendo concentrar el MVP en HanGo Learn. La ruta desde cero, Han Tutor, Han Companion, Han Vocabulary, los dos sistemas de Speaking, la personalización y HanCoins pertenecen a Learn, aunque no todos deben convertirse en puertas obligatorias para la beta. Teacher, Institutions, Connect, Community, Play y Marketplace continúan como visión de largo plazo y no deben competir con la validación del núcleo.",
        "La primera hipótesis de usuario es un estudiante hispanohablante que estudia mandarín con intención real de mejorar, necesita practicar entre clases y valora feedback claro. HanGo ya ofrece una introducción cultural y lingüística permanente para principiantes absolutos, pero la beta no debe prometer todavía una ruta completa desde cero; esa experiencia profunda se validará posteriormente como HanGo Learn 2.0.":
            "HanGo se dirige a estudiantes hispanohablantes desde cero y también a quienes ya estudian mandarín y quieren perfeccionarlo. La primera cohorte deberá acotar un tramo concreto de esa promesa para poder medirlo bien, pero el inicio desde cero ya forma parte de la dirección actual del producto y no se presenta como una línea futura separada.",
        "HanGo Learn es el producto que ya se está construyendo y el único que debe demostrar valor en el MVP. Su función es acompañar al estudiante desde una situación de aprendizaje hasta la práctica, el feedback y el siguiente paso recomendado.":
            "HanGo Learn es el producto que ya se está construyendo y el único que debe demostrar valor en el MVP. Su función es acompañar al estudiante desde los fundamentos iniciales hasta situaciones reales, práctica, feedback y un siguiente paso recomendado de acuerdo con su nivel.",
        "Han Tutor y Han Practice son experiencias diferentes: una explica; la otra practica mandarín.":
            "Han Tutor lateral explica y responde con pedagogía adaptativa. Su modo de práctica sigue siendo distinto de Han Companion, que ofrece una conversación personal por voz.",
        "El Speaking de lección representa un personaje y objetivo; Han es el profesor lateral de la aplicación.":
            "El Speaking de Learning Journey representa un personaje y un objetivo de lección; Speaking Lab es una práctica independiente. Ninguno de los dos debe compartir identidad, memoria ni propósito con Han Tutor o Han Companion.",
        "Logros, celebraciones y leaderboard semanal como hipótesis de retención; deben auditarse para evitar que la competencia desplace el aprendizaje.":
            "Logros, celebraciones y leaderboard acumulativo por puntos de Nivel General. Un premio semanal separado entrega HanCoins según lo ganado durante la semana; ambos deben auditarse para evitar que la competencia desplace el aprendizaje.",
        "Ruta piloto de 6-8 lecciones por nivel con dificultad coherente y objetivos claros.":
            "Ruta piloto que permita comenzar desde cero y avanzar hacia situaciones reales, limitada a 6-8 lecciones coherentes y objetivos medibles para la primera cohorte.",
        "Han Tutor y Han Practice estables, diferenciados y adaptados al nivel.":
            "Han Tutor estable, natural y adaptado al nivel, con ejemplos en caracteres, pinyin y traducción; Han Companion se mantiene como experiencia de voz separada.",
        "Writing básico funcional; evaluación por foto solo si calidad y costo son aceptables. Mencionar que aún no se ha utilizado.":
            "Writing básico funcional. La evaluación manuscrita por foto está construida, pero aún no se ha validado con usuarios y solo entra si su calidad y costo son aceptables.",
        "Despliegue, monitoreo, backups, privacidad, límites de gasto y soporte de beta.":
            "Despliegue, monitoreo, backups, privacidad, registro de tokens y costo por función, rate limits, reintentos idempotentes y soporte de beta.",
        "Experiencia específica para mandarín y usuarios hispanohablantes.":
            "Experiencia específica para mandarín y usuarios hispanohablantes, desde principiante absoluto hasta perfeccionamiento.",
        "Calidad: fallos graves de Han, Speaking, Writing, nivel, contenido y progreso.":
            "Calidad: fallos graves de Han Tutor, Han Companion, los dos Speakings, Writing, transcripción, nivel, contenido y progreso.",
        "Economía: costo real por usuario, por minuto de voz y por evaluación.":
            "Economía: tokens de entrada y salida, costo por usuario, minuto de voz, evaluación y función; además de reintentos y consumo evitado por caché o contenido determinista.",
        "3.  Usar contenido corto y TikTok para comunicar el problema, mostrar experiencias reales y captar testers medibles.":
            "3.  Usar la landing pública, el preview social y contenido corto para comunicar que HanGo sirve desde cero hasta perfeccionamiento, mostrar experiencias reales y captar testers medibles.",
        "Medir texto, transcripción, TTS, voz en tiempo real, visión, storage, hosting y soporte por separado.":
            "Medir texto, tokens de entrada y salida, transcripción, TTS, voz en tiempo real, visión, storage, hosting y soporte por separado mediante registros de uso por proveedor y modelo.",
        "Definir límites, rate limits y alertas antes de ofrecer uso amplio.":
            "Mantener presupuestos máximos de salida y rate limits por flujo. Las evaluaciones más costosas tienen techos explícitos; placement speaking usa hasta 1.800 tokens de salida y Reflection o Speaking Lab hasta 1.600.",
        "No confundir costo de pruebas internas con costo de un usuario activo real.":
            "No confundir el techo técnico de tokens ni el costo de pruebas internas con el consumo esperado de un usuario activo real.",
        "Ruta desde cero para principiantes absolutos: pinyin, tonos, caracteres y hábitos iniciales.":
            "Profundización adaptativa de la ruta desde cero ya iniciada: mejor secuencia de pinyin, tonos, caracteres, repaso y transferencia hacia situaciones reales.",
        "Tokens y recompensas dentro de Learn 2.0":
            "HanCoins y recompensas dentro de HanGo Learn",
        "HanCoins ya existe como sistema interno de recompensa, separado del Nivel General y sin valor financiero. Se gana por actividad verificable -lecciones, escenarios, constancia, logros y desempeño semanal- y se utiliza únicamente en personalización cosmética del avatar. Learn 2.0 podría ampliar sus usos, pero solo si la beta demuestra que mejora retorno y motivación sin crear pay-to-win, ansiedad, competencia tóxica o acumulación vacía.":
            "HanCoins ya funciona como sistema interno de recompensa, separado de los puntos de Nivel General y sin valor financiero. Se gana por actividad verificable, constancia, logros y premios semanales, y se usa únicamente en la personalización cosmética del avatar. Cualquier ampliación debe esperar evidencia de que mejora retorno y motivación sin crear pay-to-win, ansiedad, competencia tóxica o acumulación vacía.",
        "LÍMITES DE MARCA  Han es el tutor de IA dentro de HanGo Learn, no una línea separada. Speaking, Writing, Progress, los labs, los avatares y HanCoins también pertenecen a Learn. Los juegos entre alumnos pertenecen a HanGo Play. Una futura ampliación de recompensas pertenece a Learn 2.0; la creación y venta de contenido pertenece a HanGo Marketplace.":
            "LÍMITES DE MARCA  Han es el tutor de IA dentro de HanGo Learn, no una línea separada. Han Tutor, Han Companion, Han Vocabulary, Speaking, Writing, Progress, los labs, los avatares y HanCoins pertenecen a Learn. Los juegos entre alumnos pertenecen a HanGo Play. Cualquier ampliación de HanCoins se evaluará como evolución de Learn; la creación y venta de contenido pertenece a HanGo Marketplace.",
    }

    for old, new in replacements.items():
        replace_exact(doc, old, new)

    speaking_anchor = next(
        p for p in doc.paragraphs if p.text.startswith("Speaking contextual de lección")
    )
    insert_before(
        speaking_anchor,
        "Han Vocabulary como consulta contextual ajustada automáticamente al nivel, sin puntaje ni HanCoins.",
        "List Bullet",
    )
    insert_before(
        speaking_anchor,
        "Han Companion como conversación personal por voz, sin transcript visible ni mezcla con el tutor o los Speakings.",
        "List Bullet",
    )

    table2 = doc.tables[2]
    update_table_cell(table2, 1, 2, "Existe una experiencia amplia desde cero hasta perfeccionamiento; falta reducirla a una ruta piloto coherente y operable.")
    update_table_cell(table2, 3, 1, "Base curricular amplia desde fundamentos")
    update_table_cell(table2, 3, 2, "El MVP validará un tramo de 6-8 lecciones con pinyin cuando corresponda y revisión humana del mandarín.")
    update_table_cell(table2, 4, 2, "Han Tutor, Han Companion, ambos Speakings, Writing y Progress han avanzado, pero deben probarse como un sistema completo.")

    table3 = doc.tables[3]
    update_table_cell(table3, 1, 2, "Aprender mandarín desde cero y perfeccionarlo con ruta, IA, voz, escritura, vocabulario y progreso.")

    table4 = doc.tables[4]
    update_table_cell(table4, 2, 1, "Organiza fundamentos, situaciones, niveles, lecciones, acceso y continuidad desde cero hasta niveles avanzados.")
    update_table_cell(table4, 4, 1, "Explica y responde con pedagogía adaptativa; usa caracteres, pinyin y traducción en ejemplos escritos sin leer el pinyin en voz alta.")
    update_table_cell(table4, 5, 0, "Modo Practice de Han Tutor")
    update_table_cell(table4, 5, 1, "Practica mandarín dentro del tutor sin convertirse en Han Companion ni en un sistema de Speaking.")
    update_table_cell(table4, 12, 1, "Avatares zodiacales, HanCoins, tienda, logros, leaderboard acumulativo y premio semanal separado.")
    append_table_row(
        table4,
        [
            "Han Companion",
            "Conversación personal por voz con memoria de hechos compartidos; orbe con bordes distintos para alumno y Han, sin chat ni transcript visible.",
            "Construido; QA continuo de voz, silencio y memoria",
        ],
        1,
    )
    append_table_row(
        table4,
        [
            "Han Vocabulary",
            "Consulta palabras por caracteres o pinyin y devuelve significado, ejemplos, uso y errores comunes al nivel real del alumno.",
            "Construido; herramienta de referencia sin puntos",
        ],
        1,
    )
    append_table_row(
        table4,
        [
            "Landing pública",
            "Presenta HanGo en español desde cero hasta perfeccionamiento, muestra el ciclo de aprendizaje y dirige a registro o acceso.",
            "Construida; falta medir conversión",
        ],
        1,
    )

    table9 = doc.tables[9]
    update_table_cell(table9, 1, 2, "Beta productiva, ruta piloto desde cero, calidad, retorno, aprendizaje, costo y conversión de la landing.")
    update_table_cell(table9, 2, 2, "Profundizar fundamentos, Han Companion, Han Vocabulary, shadowing y roleplays según evidencia de uso.")

    table11 = doc.tables[11]
    update_table_cell(table11, 4, 2, "Registrar tokens y costo por función; aplicar techos de salida, rate limits, caché, idempotencia y alertas.")
    update_table_cell(table11, 10, 0, "Gamificación mal calibrada")
    update_table_cell(table11, 10, 1, "HanCoins, logros o ranking desplazan el aprendizaje o incentivan conductas vacías.")
    update_table_cell(table11, 10, 2, "Mantener HanCoins separados de los puntos de Nivel General y medir retorno, calidad y comportamiento.")
    append_table_row(
        table11,
        [
            "Falsos turnos o transcripciones inventadas",
            "El sistema atribuye al alumno palabras que no dijo y pierde confianza.",
            "Rechazar silencio, ruido y firmas del prompt antes de mostrar, guardar o enviar el turno; pedir repetición cuando la evidencia sea débil.",
        ],
        1,
    )
    append_table_row(
        table11,
        [
            "Sistemas de voz confundidos",
            "Han Tutor, Han Companion, Speaking de Learning Journey y Speaking Lab comparten identidad u objetivos incorrectos.",
            "Mantener rutas, prompts, memoria, métricas y pruebas separadas para cada experiencia.",
        ],
        1,
    )

    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Plan personalizado con objetivos"):
            for run in paragraph.runs:
                run.font.highlight_color = None
            paragraph_properties = paragraph._p.get_or_add_pPr()
            paragraph_run_properties = paragraph_properties.find(qn("w:rPr"))
            if paragraph_run_properties is not None:
                highlight = paragraph_run_properties.find(qn("w:highlight"))
                if highlight is not None:
                    paragraph_run_properties.remove(highlight)

    doc.core_properties.title = "HanGo Documento estratégico para socios"
    doc.core_properties.subject = "Producto actual alcance del MVP y visión de ecosistema"
    doc.core_properties.modified = datetime(2026, 9, 5, 12, 0, 0)
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
